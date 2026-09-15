import email
import html
import io
import re
from email import policy

from docx import Document
from pypdf import PdfReader


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        return _docx_text(data)
    if ext == "txt":
        return data.decode("utf-8", errors="replace")
    if ext == "eml":
        return _eml_text(data)
    raise UnsupportedDocument(f"Unsupported file type '.{ext}'. Use PDF, DOCX, TXT or EML.")


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def _eml_text(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=policy.default)
    headers = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Date", "Subject") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body else ""
    if body is not None and body.get_content_type() == "text/html":
        text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    return f"{headers}\n\n{text}"
